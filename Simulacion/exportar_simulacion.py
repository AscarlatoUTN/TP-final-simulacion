"""
Genera un Excel "pro" comparando los 5 escenarios de flota (ESCENARIOS_FLOTA)
para UNA franja horaria (FRANJA), con 5 hojas:
- Resumen: tabla comparativa de todas las métricas, una columna por escenario.
- KPIs: panel visual con los indicadores más importantes.
- Insights: texto automático con hallazgos y recomendaciones por escenario.
- Gráficos: comparaciones visuales entre escenarios.
- Eficiencia: ranking de escenarios según un score compuesto.

Ejecutalo directo con el botón Run del IDE, o con: python exportar_excel_pro.py
"""

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from config import TARIFA_BASE_FRANJA
from simulacion import Simulacion

# --- Parámetros hardcodeados: franja a analizar y escenarios de flota a comparar ---
FRANJA = "madrugada"  # una de: "madrugada", "manana", "tarde", "noche"
ESCENARIOS_FLOTA = [
    (5, 5, 5),
    (10, 10, 10),
    (15, 10, 5),
    (5, 10, 15),
    (20, 15, 10),
]
SEED = 1
ARCHIVO_SALIDA = "resultados_simulacion_pro.xlsx"

# --- Pesos del score de eficiencia compuesto (mayor score = mejor) ---
PESO_ESPERA = 0.4
PESO_OCIOSO = 0.3
PESO_PARR = 0.2
PESO_BN = 0.1

# ── Paleta ────────────────────────────────────────────────────────────────
AZUL_OSCURO = "0D2137"
AZUL_MED = "1A4A7A"
AZUL_CLARO = "2E86C1"
AZUL_SUAVE = "D6EAF8"
GRIS_CLARO = "F2F3F4"
VERDE_OK = "1E8449"
NARANJA_WARN = "CA6F1E"
ROJO_BAD = "C0392B"
BLANCO = "FFFFFF"
GOLD = "F4D03F"

COLOR_TIPO = {
    "Convencional": AZUL_CLARO,
    "Eléctrico": "117A65",
    "Autónomo": "7B2400",
}


# ── Helpers ───────────────────────────────────────────────────────────────
def _thin(color="BBBBBB"):
    s = Side(style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)


def cel(ws, ref, value=None, bold=False, size=10, color="000000",
        bg=None, align="center", border=False, fmt=None, italic=False,
        wrap=True):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = Font(name="Calibri", bold=bold, italic=italic, size=size, color=color)
    if bg:
        c.fill = PatternFill("solid", fgColor=bg)
    c.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap)
    if border:
        c.border = _thin()
    if fmt:
        c.number_format = fmt


def set_col_widths(ws, widths):
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w


def set_row_height(ws, row, height):
    ws.row_dimensions[row].height = height


def color_semaforo(valor_pct, umbral_alto=0.20, umbral_medio=0.10):
    """Rojo si supera umbral_alto, naranja si supera umbral_medio, verde si no."""
    if valor_pct > umbral_alto:
        return ROJO_BAD
    if valor_pct > umbral_medio:
        return NARANJA_WARN
    return VERDE_OK


# ── Simulación: correr los 5 escenarios para la franja elegida ────────────
def correr_escenarios():
    resultados = []
    for n_escenario, (nc, ne, na) in enumerate(ESCENARIOS_FLOTA, start=1):
        sim = Simulacion(
            cantidad_convencionales=nc,
            cantidad_electricos=ne,
            cantidad_autonomos=na,
            franja=FRANJA,
            tarifa_base=TARIFA_BASE_FRANJA[FRANJA],
            seed=SEED,
        )
        sim.correr()

        resultados.append({
            "nombre": f"Esc. {n_escenario} ({nc}C/{ne}E/{na}A)",
            "n_convencionales": nc,
            "n_electricos": ne,
            "n_autonomos": na,
            "solicitudes": sim._cantidad_solicitudes,
            "atendidos": sim._cantidad_viajes_completados,
            "arrepentidos": sim._cantidad_arrepentidos,
            "pct_arrepentidos": sim.PARR / 100,  # fracción, no 0-100
            "tiempo_espera_prom_min": sim.TPE / 60,
            "beneficio_neto": sim.BN,
            "ocioso": {
                "Convencional": sim.TPOC / 100,
                "Eléctrico": sim.TPOE / 100,
                "Autónomo": sim.TPOA / 100,
            },
            "calificacion": {
                "Convencional": sim.calificaciones.promedio("convencional"),
                "Eléctrico": sim.calificaciones.promedio("electrico"),
                "Autónomo": sim.calificaciones.promedio("autonomo"),
            },
        })
    return resultados


def _normalizar(valores, menor_es_mejor):
    """Normaliza una lista de valores a [0, 1], donde 1 siempre es 'mejor'."""
    minimo, maximo = min(valores), max(valores)
    rango = maximo - minimo
    if rango == 0:
        return [1.0 for _ in valores]
    if menor_es_mejor:
        return [(maximo - v) / rango for v in valores]
    return [(v - minimo) / rango for v in valores]


def calcular_eficiencia(resultados):
    esperas = [r["tiempo_espera_prom_min"] for r in resultados]
    ociosos_prom = [sum(r["ocioso"].values()) / 3 for r in resultados]
    parrs = [r["pct_arrepentidos"] for r in resultados]
    bns = [r["beneficio_neto"] for r in resultados]

    n_espera = _normalizar(esperas, menor_es_mejor=True)
    n_ocioso = _normalizar(ociosos_prom, menor_es_mejor=True)
    n_parr = _normalizar(parrs, menor_es_mejor=True)
    n_bn = _normalizar(bns, menor_es_mejor=False)

    for i, r in enumerate(resultados):
        r["eficiencia_global"] = (
                PESO_ESPERA * n_espera[i]
                + PESO_OCIOSO * n_ocioso[i]
                + PESO_PARR * n_parr[i]
                + PESO_BN * n_bn[i]
        )


# ── Hoja 1: Resumen comparativo ────────────────────────────────────────────
def _hoja_resumen(wb, resultados):
    ws = wb.active
    ws.title = "Resumen"
    ws.sheet_view.showGridLines = False

    n = len(resultados)
    ultima_col = get_column_letter(1 + n)

    ws.merge_cells(f"A1:{ultima_col}1")
    cel(ws, "A1", f"ANÁLISIS DE ESCENARIOS DE FLOTA — FRANJA {FRANJA.upper()}",
        bold=True, size=15, color=BLANCO, bg=AZUL_OSCURO)
    set_row_height(ws, 1, 36)

    ws.merge_cells(f"A2:{ultima_col}2")
    cel(ws, "A2",
        f"  {n} escenarios  ·  Seed: {SEED}  ·  Tarifa base: ${TARIFA_BASE_FRANJA[FRANJA]:.2f}",
        italic=True, size=9, color=BLANCO, bg=AZUL_MED, align="left")
    set_row_height(ws, 2, 18)

    ws.merge_cells(f"A3:{ultima_col}3")
    ws["A3"].fill = PatternFill("solid", fgColor=AZUL_CLARO)
    set_row_height(ws, 3, 4)

    # Fila 4: nombre de cada escenario
    set_row_height(ws, 4, 22)
    cel(ws, "A4", "Métrica", bold=True, size=10, color=BLANCO, bg=AZUL_MED, border=True)
    for i, r in enumerate(resultados):
        col = get_column_letter(2 + i)
        cel(ws, f"{col}4", r["nombre"], bold=True, size=10, color=GOLD, bg=AZUL_OSCURO, border=True)

    METRICAS = [
        ("N° convencionales", "n_convencionales", None),
        ("N° eléctricos", "n_electricos", None),
        ("N° autónomos", "n_autonomos", None),
        ("Solicitudes totales", "solicitudes", None),
        ("Viajes completados", "atendidos", None),
        ("Pasajeros arrepentidos", "arrepentidos", None),
        ("% Arrepentimiento (PARR)", "pct_arrepentidos", "0.0%"),
        ("Espera promedio (min)", "tiempo_espera_prom_min", "0.0"),
        ("Beneficio neto (BN)", "beneficio_neto", "$#,##0.00"),
    ]

    fila = 5
    for idx, (label, key, fmt) in enumerate(METRICAS):
        row_bg = GRIS_CLARO if idx % 2 == 0 else BLANCO
        set_row_height(ws, fila, 20)
        cel(ws, f"A{fila}", label, bold=True, size=10, bg=row_bg, align="left", border=True)
        for i, r in enumerate(resultados):
            col = get_column_letter(2 + i)
            val = r[key]
            txt_color = "000000"
            if key == "pct_arrepentidos":
                txt_color = color_semaforo(val)
            cel(ws, f"{col}{fila}", val, size=10, color=txt_color, bg=row_bg, border=True, fmt=fmt)
        fila += 1

    # % ocioso y calificación, por tipo
    for tipo in ("Convencional", "Eléctrico", "Autónomo"):
        row_bg = GRIS_CLARO if fila % 2 == 0 else BLANCO
        set_row_height(ws, fila, 20)
        cel(ws, f"A{fila}", f"% Ocioso {tipo}", bold=True, size=10, bg=row_bg, align="left", border=True)
        for i, r in enumerate(resultados):
            col = get_column_letter(2 + i)
            cel(ws, f"{col}{fila}", r["ocioso"][tipo], size=10, color=COLOR_TIPO[tipo],
                bg=row_bg, border=True, fmt="0.0%")
        fila += 1

    for tipo in ("Convencional", "Eléctrico", "Autónomo"):
        row_bg = GRIS_CLARO if fila % 2 == 0 else BLANCO
        set_row_height(ws, fila, 20)
        cel(ws, f"A{fila}", f"Calificación {tipo}", bold=True, size=10, bg=row_bg, align="left", border=True)
        for i, r in enumerate(resultados):
            col = get_column_letter(2 + i)
            calif = r["calificacion"][tipo]
            color_calif = VERDE_OK if calif >= 4.0 else (NARANJA_WARN if calif >= 3.5 else ROJO_BAD)
            cel(ws, f"{col}{fila}", calif, size=10, color=color_calif, bg=row_bg, border=True, fmt="0.00")
        fila += 1

    # Eficiencia global
    fila += 1
    set_row_height(ws, fila, 24)
    cel(ws, f"A{fila}", "Eficiencia global (↑ mayor = mejor)",
        bold=True, size=10, bg=AZUL_OSCURO, color=BLANCO, border=True)
    mejor_score = max(r["eficiencia_global"] for r in resultados)
    for i, r in enumerate(resultados):
        col = get_column_letter(2 + i)
        score = r["eficiencia_global"]
        es_mejor = abs(score - mejor_score) < 1e-9
        cel(ws, f"{col}{fila}", score, size=11, bold=True,
            color=GOLD if es_mejor else "AAAAAA", bg=AZUL_OSCURO, border=True, fmt="0.000")

    anchos = {"A": 28}
    for i in range(n):
        anchos[get_column_letter(2 + i)] = 16
    set_col_widths(ws, anchos)


# ── Hoja 2: KPIs ────────────────────────────────────────────────────────────
def _hoja_kpis(wb, resultados):
    ws = wb.create_sheet("KPIs")
    ws.sheet_view.showGridLines = False
    n = len(resultados)

    ws.merge_cells(f"A1:{get_column_letter(1 + n)}1")
    cel(ws, "A1", "PANEL DE KPIs POR ESCENARIO", bold=True, size=14, color=BLANCO, bg=AZUL_OSCURO)
    set_row_height(ws, 1, 32)

    ws.merge_cells(f"A2:{get_column_letter(1 + n)}2")
    cel(ws, "A2", f"Comparación rápida — franja {FRANJA}", italic=True, size=9, color=BLANCO, bg=AZUL_MED)
    set_row_height(ws, 2, 16)

    mejor_score = max(r["eficiencia_global"] for r in resultados)
    for i, r in enumerate(resultados):
        col = get_column_letter(2 + i)
        es_mejor = abs(r["eficiencia_global"] - mejor_score) < 1e-9
        bg_header = "117A65" if es_mejor else AZUL_CLARO
        cel(ws, f"{col}4", r["nombre"] + (" ★" if es_mejor else ""),
            bold=True, size=11, color=BLANCO, bg=bg_header, border=True)
    cel(ws, "A4", "", bg=AZUL_OSCURO)
    set_row_height(ws, 4, 24)

    KPIS = [
        ("% Arrepentimiento", "pct_arrepentidos", "0.0%", True),
        ("Espera prom. (min)", "tiempo_espera_prom_min", "0.0", False),
        ("Beneficio neto", "beneficio_neto", "$#,##0", False),
        ("Viajes completados", "atendidos", "0", False),
    ]

    fila = 5
    for label, key, fmt, es_pct in KPIS:
        cel(ws, f"A{fila}", label, bold=True, size=10, color=BLANCO, bg=AZUL_OSCURO, border=True)
        for i, r in enumerate(resultados):
            col = get_column_letter(2 + i)
            val = r[key]
            txt_color = color_semaforo(val) if es_pct else AZUL_CLARO
            cel(ws, f"{col}{fila}", val, bold=True, size=11, color=txt_color, bg=BLANCO, border=True, fmt=fmt)
        fila += 2

    for tipo in ("Convencional", "Eléctrico", "Autónomo"):
        cel(ws, f"A{fila}", f"% Ocioso {tipo}", bold=True, size=10, color=BLANCO, bg=AZUL_OSCURO, border=True)
        for i, r in enumerate(resultados):
            col = get_column_letter(2 + i)
            cel(ws, f"{col}{fila}", r["ocioso"][tipo], bold=True, size=11,
                color=COLOR_TIPO[tipo], bg=BLANCO, border=True, fmt="0.0%")
        fila += 2

    for f in range(5, fila, 2):
        for c in range(1, 2 + n):
            ws.cell(row=f + 1, column=c).fill = PatternFill("solid", fgColor=AZUL_SUAVE)

    anchos = {"A": 20}
    for i in range(n):
        anchos[get_column_letter(2 + i)] = 16
    set_col_widths(ws, anchos)


# ── Hoja 3: Insights ─────────────────────────────────────────────────────────
def _hoja_insights(wb, resultados):
    ws = wb.create_sheet("Insights")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:E1")
    cel(ws, "A1", "INSIGHTS AUTOMÁTICOS", bold=True, size=14, color=BLANCO, bg=AZUL_OSCURO)
    set_row_height(ws, 1, 32)

    row = 3
    for r in resultados:
        ws.merge_cells(f"A{row}:E{row}")
        cel(ws, f"A{row}", f"Escenario: {r['nombre']}", bold=True, size=12, color=BLANCO, bg=AZUL_MED, align="left")
        set_row_height(ws, row, 24)
        row += 1

        ocioso_prom = sum(r["ocioso"].values()) / 3
        lines = [
            ("% Arrepentimiento:", f"{r['pct_arrepentidos']:.1%}", color_semaforo(r["pct_arrepentidos"])),
            ("Viajes completados:", f"{r['atendidos']:.0f} / {r['solicitudes']:.0f} solicitudes", AZUL_CLARO),
            ("Espera promedio:", f"{r['tiempo_espera_prom_min']:.1f} min", "333333"),
            ("% Ocioso promedio (3 tipos):", f"{ocioso_prom:.1%}", "333333"),
            ("Beneficio neto:", f"${r['beneficio_neto']:,.2f}", VERDE_OK if r["beneficio_neto"] > 0 else ROJO_BAD),
            ("Eficiencia global:", f"{r['eficiencia_global']:.3f}  (↑ mayor = mejor)", AZUL_OSCURO),
        ]
        for label, value, vc in lines:
            set_row_height(ws, row, 18)
            cel(ws, f"A{row}", label, bold=True, size=10, bg=GRIS_CLARO, align="left", border=True)
            ws.merge_cells(f"B{row}:C{row}")
            cel(ws, f"B{row}", value, bold=True, size=10, color=vc, bg=BLANCO, align="left", border=True)
            row += 1

        set_row_height(ws, row, 20)
        ws.merge_cells(f"A{row}:E{row}")
        if r["pct_arrepentidos"] > 0.20:
            rec = "Alta pérdida de pasajeros — la flota está desbordada, considerar agregar más taxis."
            bg_rec = "FADBD8"
        elif r["pct_arrepentidos"] > 0.10:
            rec = "Pérdida moderada de pasajeros — monitorear de cerca esta composición de flota."
            bg_rec = "FDEBD0"
        else:
            rec = "Buen equilibrio entre oferta y demanda para esta franja."
            bg_rec = "D5F5E3"
        cel(ws, f"A{row}", rec, italic=True, size=10, bg=bg_rec, align="left")
        row += 2

    set_col_widths(ws, {"A": 26, "B": 22, "C": 16, "D": 14, "E": 14})


# ── Hoja 4: Gráficos comparativos ────────────────────────────────────────────
def _hoja_graficos(wb, resultados):
    ws = wb.create_sheet("Gráficos")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:N1")
    cel(ws, "A1", "GRÁFICOS COMPARATIVOS", bold=True, size=14, color=BLANCO, bg=AZUL_OSCURO)
    set_row_height(ws, 1, 32)

    headers = [
        "Escenario", "% Arrepentimiento", "Espera prom. (min)",
        "Ocioso Conv. (%)", "Ocioso Eléc. (%)", "Ocioso Autón. (%)",
        "Beneficio neto",
    ]
    for i, h in enumerate(headers):
        cel(ws, f"{get_column_letter(i + 1)}3", h, bold=True, size=9, color=BLANCO, bg=AZUL_MED, border=True)
    set_row_height(ws, 3, 24)

    for idx, r in enumerate(resultados):
        row = 4 + idx
        vals = [
            r["nombre"], r["pct_arrepentidos"], r["tiempo_espera_prom_min"],
            r["ocioso"]["Convencional"], r["ocioso"]["Eléctrico"], r["ocioso"]["Autónomo"],
            r["beneficio_neto"],
        ]
        fmt_map = [None, "0.0%", "0.0", "0.0%", "0.0%", "0.0%", "$#,##0"]
        bg = GRIS_CLARO if idx % 2 == 0 else BLANCO
        for i, v in enumerate(vals):
            cel(ws, f"{get_column_letter(i + 1)}{row}", v, size=10, bg=bg, border=True, fmt=fmt_map[i])
        set_row_height(ws, row, 18)

    n = len(resultados)

    ch1 = BarChart()
    ch1.type = "col"
    ch1.title = "% Arrepentimiento por escenario"
    ch1.style = 10
    ch1.width, ch1.height = 18, 10
    ch1.y_axis.numFmt = "0%"
    data1 = Reference(ws, min_col=2, min_row=3, max_row=3 + n)
    cats1 = Reference(ws, min_col=1, min_row=4, max_row=3 + n)
    ch1.add_data(data1, titles_from_data=True)
    ch1.set_categories(cats1)
    ch1.series[0].graphicalProperties.solidFill = ROJO_BAD
    ws.add_chart(ch1, "I3")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.title = "% Ocioso por tipo y escenario"
    ch2.grouping = "clustered"
    ch2.style = 10
    ch2.width, ch2.height = 18, 10
    ch2.y_axis.numFmt = "0%"
    data2 = Reference(ws, min_col=4, min_row=3, max_col=6, max_row=3 + n)
    cats2 = Reference(ws, min_col=1, min_row=4, max_row=3 + n)
    ch2.add_data(data2, titles_from_data=True)
    ch2.set_categories(cats2)
    ch2.series[0].graphicalProperties.solidFill = COLOR_TIPO["Convencional"]
    ch2.series[1].graphicalProperties.solidFill = COLOR_TIPO["Eléctrico"]
    ch2.series[2].graphicalProperties.solidFill = COLOR_TIPO["Autónomo"]
    ws.add_chart(ch2, "I22")

    ch3 = BarChart()
    ch3.type = "col"
    ch3.title = "Beneficio neto por escenario"
    ch3.style = 10
    ch3.width, ch3.height = 18, 10
    data3 = Reference(ws, min_col=7, min_row=3, max_row=3 + n)
    cats3 = Reference(ws, min_col=1, min_row=4, max_row=3 + n)
    ch3.add_data(data3, titles_from_data=True)
    ch3.set_categories(cats3)
    ch3.series[0].graphicalProperties.solidFill = VERDE_OK
    ws.add_chart(ch3, "I41")

    set_col_widths(ws, {
        "A": 22, "B": 16, "C": 16, "D": 14, "E": 14, "F": 14, "G": 14, "H": 2,
    })


# ── Hoja 5: Ranking de eficiencia ────────────────────────────────────────────
def _hoja_eficiencia(wb, resultados):
    ws = wb.create_sheet("Eficiencia")
    ws.sheet_view.showGridLines = False

    ws.merge_cells("A1:F1")
    cel(ws, "A1", "RANKING DE EFICIENCIA GLOBAL", bold=True, size=14, color=BLANCO, bg=AZUL_OSCURO)
    set_row_height(ws, 1, 32)

    ws.merge_cells("A2:F2")
    cel(ws, "A2",
        f"Score compuesto (↑ mayor = mejor):  "
        f"{PESO_ESPERA}×(1-Espera norm.) + {PESO_OCIOSO}×(1-Ocioso norm.) "
        f"+ {PESO_PARR}×(1-PARR norm.) + {PESO_BN}×(BN norm.)",
        italic=True, size=8, color=BLANCO, bg=AZUL_MED, align="left")
    set_row_height(ws, 2, 16)

    hdrs = ["Pos.", "Escenario", "Score", "Evaluación"]
    for i, h in enumerate(hdrs):
        cel(ws, f"{get_column_letter(i + 1)}4", h, bold=True, size=10, color=BLANCO, bg=AZUL_MED, border=True)
    set_row_height(ws, 4, 20)

    ranked = sorted(resultados, key=lambda r: r["eficiencia_global"], reverse=True)
    medallas = ["1°", "2°", "3°", "4°", "5°"]

    for rank, r in enumerate(ranked):
        row = 5 + rank
        set_row_height(ws, row, 22)
        bg = [BLANCO, GRIS_CLARO, AZUL_SUAVE][rank % 3]
        eval_txt = "Óptimo" if rank == 0 else ("Aceptable" if rank <= 2 else "Mejorable")
        eval_col = VERDE_OK if rank == 0 else (NARANJA_WARN if rank <= 2 else ROJO_BAD)

        cel(ws, f"A{row}", medallas[rank] if rank < len(medallas) else f"{rank+1}°",
            bold=True, size=12, bg=bg, border=True)
        cel(ws, f"B{row}", r["nombre"], bold=True, size=11, bg=bg, align="left", border=True)
        cel(ws, f"C{row}", r["eficiencia_global"], bold=True, size=12,
            color=AZUL_OSCURO, bg=bg, border=True, fmt="0.000")
        cel(ws, f"D{row}", eval_txt, bold=True, size=11, color=eval_col, bg=bg, border=True)

    n = len(resultados)
    aux1 = 5 + n + 2
    ws[f"A{aux1}"] = "Escenario"
    ws[f"B{aux1}"] = "Score"
    for i, r in enumerate(ranked):
        ws[f"A{aux1 + 1 + i}"] = r["nombre"]
        ws[f"B{aux1 + 1 + i}"] = r["eficiencia_global"]

    ch1 = BarChart()
    ch1.type = "bar"
    ch1.title = "Score de eficiencia (↑ mayor = mejor)"
    ch1.style = 10
    ch1.width, ch1.height = 18, 10
    data1 = Reference(ws, min_col=2, min_row=aux1, max_row=aux1 + n)
    cats1 = Reference(ws, min_col=1, min_row=aux1 + 1, max_row=aux1 + n)
    ch1.add_data(data1, titles_from_data=True)
    ch1.set_categories(cats1)
    ch1.series[0].graphicalProperties.solidFill = AZUL_CLARO
    ws.add_chart(ch1, "F4")

    set_col_widths(ws, {"A": 10, "B": 24, "C": 12, "D": 14, "E": 2})


# ── Entry point ───────────────────────────────────────────────────────────────
def exportar_resultados():
    resultados = correr_escenarios()
    calcular_eficiencia(resultados)

    wb = Workbook()
    _hoja_resumen(wb, resultados)
    _hoja_kpis(wb, resultados)
    _hoja_insights(wb, resultados)
    _hoja_graficos(wb, resultados)
    _hoja_eficiencia(wb, resultados)

    wb.save(ARCHIVO_SALIDA)
    print(f"Archivo generado: {ARCHIVO_SALIDA} ({len(resultados)} escenarios, franja {FRANJA})")


if __name__ == "__main__":
    exportar_resultados()